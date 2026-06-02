"""Apply schedule adjustments to project_tracking_plan.xlsx.

Adjustments:
1. MLE-3: 6/29-7/3 -> 7/1-7/7 (5 weekdays in 7月上)
2. VAL-1: 7/6-7/10 -> 7/1-7/7 (parallel with MLE-3)
3. MLE-4: 7/13-7/17 -> 7/16-7/22 (5 weekdays in 7月下)
4. VAL-2: 7/20-7/24 -> 7/16-7/22 (parallel with MLE-4)
5. TEST-1: 7/27-8/14 -> 8/3-8/14 (10 weekdays in 8月上)
6. DOC-1: 9/28-10/9 -> 10/1-10/9 (in 10月上)
"""

from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

XLSX = Path("docs/project_implementation_details/project_tracking_plan.xlsx")
BLUE = PatternFill(start_color="FF9DC3E6", end_color="FF9DC3E6", fill_type="solid")
GRAY = PatternFill(start_color="FFE7E6E6", end_color="FFE7E6E6", fill_type="solid")


def date_col_map(ws):
    """Map date string -> column index from row 1 of Gantt sheet."""
    m = {}
    for c in range(11, ws.max_column + 1):
        v = ws.cell(row=1, column=c).value
        if hasattr(v, "strftime"):
            m[v.date().isoformat()] = c
    return m


def daterange(start, end):
    d = start
    while d <= end:
        yield d
        d += timedelta(days=1)


def apply_gantt_row(ws, dcol, row, old_start, old_end, new_start, new_end,
                    weekdays_only):
    """Repaint a Gantt row: clear blue across old range, set blue on new range."""
    for d in daterange(old_start, old_end):
        col = dcol.get(d.isoformat())
        if col is None:
            continue
        is_weekend = d.weekday() >= 5
        ws.cell(row=row, column=col).fill = GRAY if is_weekend else PatternFill(fill_type=None)

    for d in daterange(new_start, new_end):
        col = dcol.get(d.isoformat())
        if col is None:
            continue
        is_weekend = d.weekday() >= 5
        if weekdays_only and is_weekend:
            ws.cell(row=row, column=col).fill = GRAY
        else:
            ws.cell(row=row, column=col).fill = BLUE


def main():
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["Gantt Chart"]
    dcol = date_col_map(ws)

    plans = [
        # row, task_id, old_start, old_end, new_start, new_end, weekdays_only
        (5, "MLE-3", "2026-06-29", "2026-07-03", "2026-07-01", "2026-07-07", True),
        (6, "VAL-1", "2026-07-06", "2026-07-10", "2026-07-01", "2026-07-07", True),
        (7, "MLE-4", "2026-07-13", "2026-07-17", "2026-07-16", "2026-07-22", True),
        (8, "VAL-2", "2026-07-20", "2026-07-24", "2026-07-16", "2026-07-22", True),
        (9, "TEST-1", "2026-07-27", "2026-08-14", "2026-08-03", "2026-08-14", False),
        (13, "DOC-1", "2026-09-28", "2026-10-09", "2026-10-01", "2026-10-09", False),
    ]

    for row, tid, os_, oe_, ns_, ne_, wko in plans:
        ws.cell(row=row, column=4).value = datetime.fromisoformat(ns_)
        ws.cell(row=row, column=5).value = datetime.fromisoformat(ne_)
        apply_gantt_row(
            ws, dcol, row,
            datetime.fromisoformat(os_).date(),
            datetime.fromisoformat(oe_).date(),
            datetime.fromisoformat(ns_).date(),
            datetime.fromisoformat(ne_).date(),
            wko,
        )
        print(f"Gantt {tid} (row {row}): {os_}~{oe_} -> {ns_}~{ne_}")

    # Update Overview sheet phases
    ov = wb["Overview"]
    ov.cell(row=5, column=2).value = "7/1 - 7/22"   # Phase 3 (filtered+VAL+append parallel)
    ov.cell(row=6, column=2).value = "8/3 - 8/14"   # Phase 4 (TEST)
    ov.cell(row=8, column=2).value = "10/1 - 10/15"  # Phase 6 (DOC-1 + FINAL)
    print("Overview phase 3/4/6 dates updated")

    # Update Daily Plan: shift MLE-3, VAL-1, MLE-4, VAL-2 dates; trim TEST-1 / DOC-1
    dp = wb["Daily Plan"]

    def set_date(row, iso):
        dp.cell(row=row, column=1).value = datetime.fromisoformat(iso)

    # MLE-3 rows 40-44 -> 7/1, 7/2, 7/3, 7/6, 7/7
    for r, d in zip([40, 41, 42, 43, 44],
                    ["2026-07-01", "2026-07-02", "2026-07-03", "2026-07-06", "2026-07-07"]):
        set_date(r, d)
    # VAL-1 rows 45-49 -> same dates (parallel)
    for r, d in zip([45, 46, 47, 48, 49],
                    ["2026-07-01", "2026-07-02", "2026-07-03", "2026-07-06", "2026-07-07"]):
        set_date(r, d)
    # MLE-4 rows 50-54 -> 7/16, 7/17, 7/20, 7/21, 7/22
    for r, d in zip([50, 51, 52, 53, 54],
                    ["2026-07-16", "2026-07-17", "2026-07-20", "2026-07-21", "2026-07-22"]):
        set_date(r, d)
    # VAL-2 rows 55-58 -> 7/16, 7/17, 7/20, 7/21 (4 rows, parallel)
    for r, d in zip([55, 56, 57, 58],
                    ["2026-07-16", "2026-07-17", "2026-07-20", "2026-07-21"]):
        set_date(r, d)
    print("Daily Plan: MLE-3/VAL-1/MLE-4/VAL-2 dates shifted (parallel pairs)")

    # TEST-1: drop rows 60-64 (the 5 prep tasks 7/27-7/31), keep 65-74 -> 8/3-8/14
    # Daily Plan structure already has rows 65-74 spanning 8/3-8/14. Just delete rows 60-64.
    dp.delete_rows(60, 5)
    print("Daily Plan: deleted rows 60-64 (TEST-1 prep, now buffer 7/23-7/31)")

    # After deletion, DOC-1 rows shift up by 5: was 105-114, now 100-109
    # Drop 3 buffer rows (was 108-110, now 103-105). After delete, remaining 7 rows
    # were originally [105,106,107,111,112,113,114] -> now [100,101,102,103,104,105,106]
    dp.delete_rows(103, 3)
    print("Daily Plan: deleted 3 DOC-1 buffer rows")

    # Now reassign DOC-1 dates. After two deletions, the 7 remaining DOC-1 rows
    # are 100-106. Map to 10/1, 10/2, 10/5, 10/6, 10/7, 10/8, 10/9.
    doc1_rows = [100, 101, 102, 103, 104, 105, 106]
    doc1_dates = ["2026-10-01", "2026-10-02", "2026-10-05", "2026-10-06",
                  "2026-10-07", "2026-10-08", "2026-10-09"]
    for r, d in zip(doc1_rows, doc1_dates):
        # Verify
        cur_id = dp.cell(row=r, column=2).value
        if cur_id != "DOC-1":
            print(f"WARN: row {r} is {cur_id} not DOC-1 — skipping")
            continue
        set_date(r, d)
    print("Daily Plan: DOC-1 dates shifted to 10/1-10/9")

    wb.save(XLSX)
    print(f"Saved {XLSX}")


if __name__ == "__main__":
    main()
