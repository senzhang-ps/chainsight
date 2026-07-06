from __future__ import annotations

from pathlib import Path
import os

from databricks import sql as dbsql
from dotenv import load_dotenv
from openpyxl import load_workbook
import pandas as pd


PROJECT_DIR = Path(__file__).resolve().parents[1]
SCENARIO_NAME = "baseline-current-params-should-be-rccp-otd"
CONFIG_WORKBOOK_FILE = (
    PROJECT_DIR
    / "scenarios"
    / SCENARIO_NAME
    / "config"
    / f"{SCENARIO_NAME}.xlsx"
)
OUTPUT_DIR = PROJECT_DIR / "data"
OUTPUT_FILE = OUTPUT_DIR / "ps_psc_sku_master_category_en.csv"
TABLE_FQN = "cdl_ps_hana_prd.sl.ps_psc_sku_master"
QUERY_COLUMNS = ["material_num", "category_en", "category", "bu_attr"]
CHUNK_SIZE = 500


def normalize_material(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value).strip()

    text = str(value).strip()
    if not text:
        return None
    if text.endswith(".0"):
        text = text[:-2]
    return text


def load_materials_from_workbook(path: Path) -> list[str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["M6_MaterialMD"]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    material_index = list(header).index("material")

    materials: set[str] = set()
    for row in rows:
        material = normalize_material(row[material_index])
        if material:
            materials.add(material)
    return sorted(materials)


def chunked(items: list[str], size: int) -> list[list[str]]:
    return [items[index:index + size] for index in range(0, len(items), size)]


def fetch_mapping(materials: list[str]) -> pd.DataFrame:
    load_dotenv(PROJECT_DIR.parents[1] / ".env")
    host = os.environ["DATABRICKS_HOST"]
    http_path = os.environ["DATABRICKS_HTTP_PATH"]
    token = os.environ["DATABRICKS_TOKEN"]

    frames: list[pd.DataFrame] = []
    with dbsql.connect(server_hostname=host, http_path=http_path, access_token=token) as conn:
        for batch in chunked(materials, CHUNK_SIZE):
            quoted = ", ".join(f"'{material}'" for material in batch)
            query = f"""
            select {', '.join(QUERY_COLUMNS)}
            from {TABLE_FQN}
            where material_num in ({quoted})
            """
            with conn.cursor() as cursor:
                cursor.execute(query)
                rows = cursor.fetchall()
                frames.append(pd.DataFrame(rows, columns=QUERY_COLUMNS))

    if not frames:
        return pd.DataFrame(columns=QUERY_COLUMNS)
    return pd.concat(frames, ignore_index=True).drop_duplicates(subset=["material_num"])


def main() -> None:
    materials = load_materials_from_workbook(CONFIG_WORKBOOK_FILE)
    mapping = fetch_mapping(materials)
    output = pd.DataFrame({"material": materials}).merge(
        mapping.rename(columns={"material_num": "material"}),
        on="material",
        how="left",
    )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output.to_csv(OUTPUT_FILE, index=False)

    matched = int(output["category_en"].notna().sum()) if "category_en" in output.columns else 0
    print(f"Material rows written: {len(output)}")
    print(f"Matched category_en rows: {matched}")
    print(f"Output written: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()