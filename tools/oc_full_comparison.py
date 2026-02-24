#!/usr/bin/env python3
"""
OC配置 三版本全量数据对比工具 (行数对比版v2)
"""
import json, warnings, sys, time
from pathlib import Path
from datetime import datetime, timedelta
import openpyxl, psycopg

warnings.filterwarnings("ignore")

PROJECT_ROOT = Path(__file__).parent.parent
DEV_DIR = PROJECT_ROOT / "outputs" / "run_20260127_142402"
SRC_DIR = PROJECT_ROOT / "outputs" / "OC_Paste_S1_20251224" / "run_20260209_222302"
DB_CONN = "host=localhost port=5432 dbname=test_db user=postgres password=123456"

START_DATE = datetime(2025, 12, 15)
NUM_DAYS = 76
SIM_DATES = [(START_DATE + timedelta(days=i)).strftime("%Y%m%d") for i in range(NUM_DAYS)]

MODULES = [
    ("module1", "module1_output_{date}.xlsx", "OrderLog",       "module1_output_orderlog",       "sim_date"),
    ("module1", "module1_output_{date}.xlsx", "ShipmentLog",    "module1_output_shipmentlog",    "sim_date"),
    ("module1", "module1_output_{date}.xlsx", "CutLog",         "module1_output_cutlog",         "sim_date"),
    ("module3", "Module3Output_{date}.xlsx",  "NetDemand",      "module3_output_netdemand",      "sim_date"),
    ("module4", "Module4Output_{date}.xlsx",  "ProductionPlan", "module4_output_productionplan", "sim_date"),
    ("module5", "Module5Output_{date}.xlsx",  "DeploymentPlan", "module5_output_deploymentplan", "sim_date"),
    ("module6", "Module6Output_{date}.xlsx",  "DeliveryPlan",   "module6_output_deliveryplan",   "sim_date"),
]

def xlsx_row_count(filepath, sheet_name):
    try:
        if not filepath.exists(): return 0
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        if sheet_name not in wb.sheetnames:
            wb.close(); return 0
        ws = wb[sheet_name]
        r = ws.max_row
        wb.close()
        return max((r or 1) - 1, 0)
    except:
        return -1

def db_batch_counts(conn, table, date_col, dates):
    result = {}
    try:
        cur = conn.cursor()
        phs = ",".join(f"'{d}'" for d in dates)
        cur.execute(f"SELECT {date_col}, COUNT(*) FROM {table} WHERE {date_col} IN ({phs}) GROUP BY {date_col}")
        for row in cur.fetchall():
            result[str(row[0])] = row[1]
        cur.close()
    except Exception as e:
        print(f"  DB ERR {table}: {e}", flush=True)
    return result

def main():
    t_start = time.time()
    print("=" * 90, flush=True)
    print("ChainSight OC 三版本全量数据一致性验证", flush=True)
    print(f"日期范围: {SIM_DATES[0]} ~ {SIM_DATES[-1]} ({NUM_DAYS}天)", flush=True)
    print("=" * 90, flush=True)

    conn = None
    try:
        conn = psycopg.connect(DB_CONN)
        print("DB连接成功", flush=True)
    except Exception as e:
        print(f"DB连接失败: {e}", flush=True)

    results = {}
    for module_dir, file_pat, sheet, db_table, db_dcol in MODULES:
        key = f"{module_dir}/{sheet}"
        t0 = time.time()
        sys.stdout.write(f"对比 {key:<30} "); sys.stdout.flush()

        db_counts = db_batch_counts(conn, db_table, db_dcol, SIM_DATES) if conn else {}

        daily = []
        td, ts, tdb = 0, 0, 0
        all_match = True
        for ds in SIM_DATES:
            fn = file_pat.format(date=ds)
            dr = xlsx_row_count(DEV_DIR / module_dir / fn, sheet)
            sr = xlsx_row_count(SRC_DIR / module_dir / fn, sheet)
            dbr = db_counts.get(ds, 0)
            td += dr; ts += sr; tdb += dbr
            m = (dr == sr == dbr)
            if not m: all_match = False
            daily.append({"date": ds, "dev": dr, "src": sr, "db": dbr, "match": m})

        st = "PASS" if all_match else "FAIL"
        print(f"[{st}] Dev={td:>10,} Src={ts:>10,} DB={tdb:>10,} ({time.time()-t0:.1f}s)", flush=True)

        if not all_match:
            for d in daily:
                if not d["match"]:
                    print(f"  DIFF {d['date']}: Dev={d['dev']} Src={d['src']} DB={d['db']}", flush=True)

        results[key] = {"total_dev": td, "total_src": ts, "total_db": tdb, "all_match": all_match, "daily": daily}

    if conn: conn.close()

    # 汇总表
    print("\n" + "=" * 100, flush=True)
    print("Section 3.1 总体验证结果", flush=True)
    print("=" * 100, flush=True)
    print(f"{'模块输出':<30} {'Dev行数':>12} {'Src行数':>12} {'DB行数':>12} {'Dev=Src':>8} {'Dev=DB':>7} {'状态':>6}", flush=True)
    print("-" * 95, flush=True)
    for k, r in results.items():
        ds = "Y" if r["total_dev"] == r["total_src"] else "N"
        dd = "Y" if r["total_dev"] == r["total_db"] else "N"
        s = "PASS" if r["all_match"] else "FAIL"
        print(f"{k:<30} {r['total_dev']:>12,} {r['total_src']:>12,} {r['total_db']:>12,} {ds:>8} {dd:>7} {s:>6}", flush=True)

    # 每日部署计划
    print("\n" + "=" * 80, flush=True)
    print("Section 3.3 每日部署计划行数 (Module5/DeploymentPlan)", flush=True)
    print("=" * 80, flush=True)
    if "module5/DeploymentPlan" in results:
        print(f"{'日期':<12} {'Dev':>10} {'Src':>10} {'DB':>10} {'一致':>6}", flush=True)
        print("-" * 55, flush=True)
        for d in results["module5/DeploymentPlan"]["daily"]:
            print(f"{d['date']:<12} {d['dev']:>10,} {d['src']:>10,} {d['db']:>10,} {'Y' if d['match'] else 'N':>6}", flush=True)

    # 每日交付计划
    print("\n" + "=" * 80, flush=True)
    print("Section 3.4 每日交付计划行数 (Module6/DeliveryPlan)", flush=True)
    print("=" * 80, flush=True)
    if "module6/DeliveryPlan" in results:
        print(f"{'日期':<12} {'Dev':>10} {'Src':>10} {'DB':>10} {'一致':>6}", flush=True)
        print("-" * 55, flush=True)
        for d in results["module6/DeliveryPlan"]["daily"]:
            print(f"{d['date']:<12} {d['dev']:>10,} {d['src']:>10,} {d['db']:>10,} {'Y' if d['match'] else 'N':>6}", flush=True)

    # 保存JSON
    out = PROJECT_ROOT / "tools" / "oc_comparison_results.json"
    with open(out, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2, default=str)
    print(f"\n结果已保存: {out}", flush=True)
    print(f"总耗时: {time.time()-t_start:.1f}s", flush=True)

if __name__ == "__main__":
    main()
